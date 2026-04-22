"""Exception workbook for the billing team.

One sheet lists every discrepancy with its cause and amount; the others give
the totals by cause and by client and the findings on the fee schedule and the
client master that sit behind them.
"""

from __future__ import annotations

import datetime as dt
import re
import zipfile
from decimal import Decimal
from pathlib import Path

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"
OVER = "F8D7D2"      # client paid too much
UNDER = "FFF2CC"     # client paid too little
UNVERIFIED = "E7E6E6"
MONEY = "#,##0.00;[Red]-#,##0.00"

# cause code -> (label, what it means)
CAUSES = {
    "DUPLICATE_CLIENT": ("Duplicate client records",
                         "The client exists twice in the master. Each record was banded on its own, so less of the volume reached the cheaper bands."),
    "EXPIRED_TERMS": ("Expired terms still applied",
                      "A negotiated discount or flat rate was charged after its end date."),
    "TERMS_NOT_APPLIED": ("Negotiated terms not applied",
                          "Terms in force for the month are missing from the invoice, which is at standard rates."),
    "OLD_RATE_CARD": ("Previous rate card",
                      "The amount matches the schedule version that had already been replaced."),
    "SLAB_INSTEAD_OF_GRADUATED": ("Slab instead of graduated bands",
                                  "The whole volume was charged at the rate of the band it reached."),
    "VOLUME_MISMATCH": ("Invoice quantity differs from activity",
                        "The amount is right for the invoiced quantity, but the quantity is not what the activity feed shows."),
    "NOT_BILLED": ("Activity not billed", "There is billable activity and no invoice line."),
    "NO_ACTIVITY": ("Billed without activity", "There is an invoice line and no activity in the month."),
    "NOT_PRICEABLE": ("Cannot be verified",
                      "The fee schedule has a defect for this service and month, so there is no price to compare with."),
    "UNEXPLAINED": ("Unexplained", "The amount matches none of the pricing variants tested."),
}

EXCEPTIONS_SQL = """
    select client_key, client_name, lei, billing_month, service_name, activity_quantity, billed_quantity,
           billed_amount, expected_amount, difference, direction, cause, explanation, invoice_nos
    from recon.v_exceptions
    order by abs(difference) desc nulls last, billed_amount desc nulls last, client_key, billing_month, service_code
"""
EXCEPTION_COLUMNS = ["Client ID", "Client", "LEI", "Month", "Service", "Activity quantity", "Invoiced quantity",
                     "Billed", "Expected", "Difference", "Direction", "Cause", "Explanation", "Invoice"]


def _plain(value):
    return float(value) if isinstance(value, Decimal) else value


def _fetch(conn: psycopg.Connection, sql: str) -> list[tuple]:
    return [tuple(_plain(v) for v in row) for row in conn.execute(sql).fetchall()]


def _header(ws, row: int, titles: list[str]) -> None:
    for column, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _table(ws, top: int, titles: list[str], rows: list[tuple], money: tuple[str, ...] = (), widths: dict | None = None) -> int:
    """Write a header and rows starting at `top`. Returns the last row used."""
    _header(ws, top, titles)
    for offset, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row=top + offset, column=column, value=value)
            if titles[column - 1] in money:
                cell.number_format = MONEY
    for column, title in enumerate(titles, start=1):
        longest = max([len(title)] + [len(str(r[column - 1])) for r in rows if r[column - 1] is not None])
        width = (widths or {}).get(title, min(max(10, longest + 2), 60))
        ws.column_dimensions[get_column_letter(column)].width = width
    return top + len(rows)


def build_workbook(conn: psycopg.Connection, path: Path) -> dict[str, int]:
    """Write the workbook and return the row count per sheet."""
    overview = _fetch(conn, "select * from recon.v_overview")[0]
    first_month, last_month = _fetch(conn, "select min(billing_month), max(billing_month) from recon.lines")[0]
    tolerance = _fetch(conn, "select amount_tolerance from recon.settings")[0][0]
    by_cause = _fetch(conn, """
        select cause, lines, over_billed, under_billed, net_difference, billed_not_verified
        from recon.v_summary_by_cause
        order by over_billed + under_billed + billed_not_verified desc, cause""")
    exceptions = _fetch(conn, EXCEPTIONS_SQL)
    by_client = _fetch(conn, """
        select client_key, client_name, exception_lines, over_billed, under_billed, net_difference
        from recon.v_client_impact
        order by over_billed + under_billed desc, client_key""")
    schedule_issues = _fetch(conn, """
        select severity, issue_code, service_code, effective_from, affected_from, affected_to, detail
        from dq.schedule_issues order by severity, service_code, effective_from, issue_code""")
    client_issues = _fetch(conn, """
        select issue_code, client_id, client_name, lei, related_client_id, detail
        from dq.client_issues order by issue_code, client_id""")

    wb = Workbook()

    # ---- Summary
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Fee billing reconciliation"
    summary["A1"].font = Font(size=16, bold=True, color=NAVY)
    summary["A2"] = (f"Billing months {first_month:%b %Y} to {last_month:%b %Y}. Difference = billed minus expected: "
                     f"positive means the client was over-billed. Differences up to {tolerance:.2f} EUR are treated as rounding.")
    lines, matched, within, mismatched, not_billed, no_activity, not_priceable, billed, _expected, over, under = overview
    facts = [
        ("Client-month-service lines compared", lines, "#,##0"),
        ("Matched exactly", matched, "#,##0"),
        ("Within tolerance", within, "#,##0"),
        ("Exceptions", mismatched + not_billed + no_activity + not_priceable, "#,##0"),
        ("Total billed (EUR)", billed, MONEY),
        ("Over-billed (EUR)", over, MONEY),
        ("Under-billed (EUR)", under, MONEY),
        ("Billed but not verifiable (EUR)", sum(row[5] for row in by_cause), MONEY),
    ]
    for offset, (label, value, number_format) in enumerate(facts):
        summary.cell(row=4 + offset, column=1, value=label)
        cell = summary.cell(row=4 + offset, column=2, value=value)
        cell.number_format = number_format
        cell.font = Font(bold=True)

    top = 4 + len(facts) + 2
    summary.cell(row=top - 1, column=1, value="By cause").font = Font(bold=True, size=12)
    cause_rows = [(CAUSES[cause][0], count, over_, under_, net, unverified)
                  for cause, count, over_, under_, net, unverified in by_cause]
    money_columns = ("Over-billed", "Under-billed", "Net difference", "Billed, not verified")
    last = _table(summary, top, ["Cause", "Lines", *money_columns], cause_rows, money=money_columns)
    total_row = last + 1
    summary.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    for column in range(2, 7):
        letter = get_column_letter(column)
        cell = summary.cell(row=total_row, column=column, value=f"=SUM({letter}{top + 1}:{letter}{last})")
        cell.font = Font(bold=True)
        cell.number_format = "#,##0" if column == 2 else MONEY
    summary.column_dimensions["A"].width = 42
    for letter in "BCDEF":
        summary.column_dimensions[letter].width = 20

    # ---- Exceptions
    sheet = wb.create_sheet("Exceptions")
    rows = [(*row[:3], f"{row[3]}", *row[4:11], CAUSES[row[11]][0], *row[12:]) for row in exceptions]
    last = _table(sheet, 1, EXCEPTION_COLUMNS, rows, money=("Billed", "Expected", "Difference"),
                  widths={"Explanation": 90, "Client": 34, "Service": 30, "Invoice": 26})
    direction_column = EXCEPTION_COLUMNS.index("Direction") + 1
    difference_column = EXCEPTION_COLUMNS.index("Difference") + 1
    fills = {"over-billed": OVER, "under-billed": UNDER, "not verified": UNVERIFIED}
    for row in range(2, last + 1):
        fill = PatternFill("solid", fgColor=fills[sheet.cell(row=row, column=direction_column).value])
        sheet.cell(row=row, column=direction_column).fill = fill
        sheet.cell(row=row, column=difference_column).fill = fill
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXCEPTION_COLUMNS))}{last}"

    # ---- By client
    sheet = wb.create_sheet("By client")
    money_columns = ("Over-billed", "Under-billed", "Net difference")
    _table(sheet, 1, ["Client ID", "Client", "Exception lines", *money_columns], by_client, money=money_columns)
    sheet.freeze_panes = "A2"

    # ---- findings behind the numbers
    sheet = wb.create_sheet("Fee schedule issues")
    _table(sheet, 1, ["Severity", "Issue", "Service", "Version from", "Affects from", "Affects to", "Detail"], schedule_issues,
           widths={"Detail": 90})
    sheet = wb.create_sheet("Client master issues")
    _table(sheet, 1, ["Issue", "Client ID", "Client", "LEI as entered", "Related client", "Detail"], client_issues,
           widths={"Detail": 90, "Client": 40})
    sheet.freeze_panes = "A2"

    sheet = wb.create_sheet("Causes")
    _table(sheet, 1, ["Cause", "What it means"], [CAUSES[code] for code in CAUSES], widths={"What it means": 120})

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, dt.date):
                    cell.number_format = "yyyy-mm-dd"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _stabilise(path, last_month)
    return {"Exceptions": len(exceptions), "By client": len(by_client),
            "Fee schedule issues": len(schedule_issues), "Client master issues": len(client_issues)}


def _stabilise(path: Path, as_of: dt.date) -> None:
    """Give the file fixed timestamps so rebuilding from the same data gives the same bytes."""
    stamp = f"{as_of:%Y-%m-%d}T00:00:00Z"
    fixed_time = (as_of.year, as_of.month, as_of.day, 0, 0, 0)
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in entries:
            if name == "docProps/core.xml":
                text = content.decode("utf-8")
                text = re.sub(r"(<dcterms:(created|modified)[^>]*>).*?(</dcterms:\2>)", rf"\g<1>{stamp}\g<3>", text)
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator>feerecon</dc:creator>", text)
                content = text.encode("utf-8")
            entry = zipfile.ZipInfo(name, date_time=fixed_time)
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, content)
