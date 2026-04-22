from __future__ import annotations

import pytest

pytest.importorskip("openpyxl")

from openpyxl import load_workbook  # noqa: E402

from feerecon.workbook import CAUSES, EXCEPTION_COLUMNS, build_workbook  # noqa: E402


@pytest.fixture(scope="module")
def workbook(built, tmp_path_factory):
    path = tmp_path_factory.mktemp("report") / "exceptions.xlsx"
    build_workbook(built, path)
    return path


def test_every_discrepancy_is_listed_with_cause_and_amount(built, workbook):
    sheet = load_workbook(workbook)["Exceptions"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    header = [cell.value for cell in sheet[1]]

    assert header == EXCEPTION_COLUMNS
    assert len(rows) == built.execute("select count(*) from recon.v_exceptions").fetchone()[0]

    cause, explanation, difference, direction = (header.index(name) for name in ("Cause", "Explanation", "Difference", "Direction"))
    labels = {label for label, _ in CAUSES.values()}
    for row in rows:
        assert row[cause] in labels and row[explanation]
        assert (row[difference] is None) == (row[direction] == "not verified")


def test_planted_error_appears_with_its_explanation(built, planted, workbook):
    client, month, service = planted["expired_discount"][0]
    sheet = load_workbook(workbook)["Exceptions"]
    match = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] == client and row[3] == month
             and row[4].startswith("Settlement")]

    assert service == "SETTLE" and len(match) == 1
    assert match[0][11] == "Expired terms still applied"
    assert "ended on" in match[0][12]
    assert match[0][10] == "under-billed" and match[0][9] < 0


def test_summary_totals_match_the_database(built, workbook):
    summary = load_workbook(workbook)["Summary"]
    facts = {row[0]: row[1] for row in summary.iter_rows(min_row=4, max_row=11, values_only=True)}
    overview = built.execute("select lines, billed_total, over_billed, under_billed from recon.v_overview").fetchone()

    assert facts["Client-month-service lines compared"] == overview[0]
    assert facts["Total billed (EUR)"] == pytest.approx(float(overview[1]))
    assert facts["Over-billed (EUR)"] == pytest.approx(float(overview[2]))
    assert facts["Under-billed (EUR)"] == pytest.approx(float(overview[3]))


def test_rebuilding_gives_the_same_file(built, workbook, tmp_path):
    again = tmp_path / "again.xlsx"
    build_workbook(built, again)

    assert again.read_bytes() == workbook.read_bytes()
